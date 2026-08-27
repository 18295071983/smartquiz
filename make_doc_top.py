# -*- coding: utf-8 -*-
"""构造测试文件：把说明文字插入题目 sheet 顶部（表头之前），验证检测/解析"""
import openpyxl

src = r'D:\qzq\smartquiz\test_files\review_full.xlsx'
dst = r'D:\qzq\smartquiz\test_files\review_doc_top.xlsx'

wb = openpyxl.load_workbook(src, data_only=True)
ws = wb['气化（1-24）']

# 读取原表所有行（含表头）
rows = []
for row in ws.iter_rows(values_only=True):
    rows.append(list(row))

# 新表：2 行说明 + 原表头 + 数据
wb2 = openpyxl.Workbook()
ws2 = wb2.active
ws2.title = '气化（1-24）'

# 说明行（模拟"模板说明"内容，置于表头之前）
doc1 = ['试题及试卷模板说明', '请按以下要求填写本表，题目录入时注意格式规范：',
        '1.题型必须填写完整（填空题/简答题/问答题等）；2.难度分为易/较易/较难/难；3.答案按题型填写：填空题填写答案内容，简答/问答题填写要点；4.题干中填空空位用（）标注。', '', '', '', '']
doc2 = ['全员综合复审题库（气化方向）', '覆盖气化工段全流程，含设备、工艺、安全知识点。', '', '', '', '', '']
ws2.append(doc1)
ws2.append(doc2)
for r in rows:
    ws2.append(r)

wb2.save(dst)
print('saved:', dst)

# 校验
wb3 = openpyxl.load_workbook(dst, read_only=True, data_only=True)
w = wb3['气化（1-24）']
cnt = 0
for row in w.iter_rows(values_only=True):
    if any(c is not None and str(c).strip() != '' for c in row):
        cnt += 1
print('非空行数:', cnt)
wb3.close()
